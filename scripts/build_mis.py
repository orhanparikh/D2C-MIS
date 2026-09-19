#!/usr/bin/env python3
"""Generate a D2C MIS workbook from a config JSON.

Usage:
  python build_mis.py config.json output.xlsx [--preserve existing.xlsx]

Config (JSON):
{
  "brand": "Acme Skincare",
  "archetype": "bpc" | "fnb" | "apparel" | "other",
  "unit": "INR lakh" | "INR" | "INR crore",
  "start_month": "2025-04",
  "months": 12,
  "channels": [{"name": "Website", "type": "website"}, {"name": "Amazon", "type": "amazon"}],
  "decisions": {"qcom_promo": "deduction", "ship_fee": "cost", "store_staff": "cm2", "exchanges": "gross"},
  "gst_overrides": {"Website": 0.18}
}

--preserve copies every input value from an existing workbook by (sheet, label, month) so a brand can add
a channel or tab without retyping. Formulas are never copied; they are regenerated.
"""
import json, sys, datetime, re
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, __import__("os").path.dirname(__file__))
from mis_spec import CHANNEL_TYPES, ARCHETYPES, DECISIONS, CHANNEL_LINES, CONSOLIDATED_SUM_KEYS, CONSOLIDATED_LINES, CASH_LINES, starter_questions

FONT = "Arial"
F_INPUT = Font(name=FONT, color="0000FF", size=10)
F_TEXT = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, bold=True, size=10)
F_LINK = Font(name=FONT, color="008000", size=10)
F_HDR = Font(name=FONT, bold=True, color="FFFFFF", size=10)
F_TITLE = Font(name=FONT, bold=True, size=14)
F_HINT = Font(name=FONT, italic=True, color="7F7F7F", size=9)
FILL_INPUT = PatternFill("solid", fgColor="FFF9E5")
FILL_HDR = PatternFill("solid", fgColor="1F3864")
FILL_SUB = PatternFill("solid", fgColor="DDEBF7")
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

START = "Start Here"
CONS = "Consolidated P&L"
CASH = "Cash & Working Capital"
CHECKS = "Checks"
DEFS = "Definitions"

FIRST_COL = 3  # C


def money_fmt(unit):
    return '#,##0.00;(#,##0.00);"-"' if unit != "INR" else '#,##0;(#,##0);"-"'


FMT = {"pct": '0.0%;(0.0%);"-"', "int": '#,##0;(#,##0);"-"', "mult": '0.0"x";(0.0"x");"-"', "days": '0;(0);"-"'}


def month_labels(start, n):
    y, m = map(int, start.split("-"))
    out = []
    for i in range(n):
        d = datetime.date(y + (m - 1 + i) // 12, (m - 1 + i) % 12 + 1, 1)
        out.append(d.strftime("%b-%y"))
    return out


def resolve(label, ct):
    return label(ct) if callable(label) else label


class Sheet:
    """Lays out lines top-down, records row per key, resolves formulas after layout."""

    def __init__(self, wb, title, cfg, months, first_row=4):
        self.ws = wb.create_sheet(title)
        self.title = title
        self.cfg = cfg
        self.months = months
        self.ncols = len(months)
        self.rows = {}
        self.row = first_row
        self.pending = []
        self.inputs = []  # (row, label) for preserve
        self.defs = []    # for Definitions tab
        self.mfmt = money_fmt(cfg["unit"])

    def col(self, i):
        return L(FIRST_COL + i)

    @property
    def total_col(self):
        return L(FIRST_COL + self.ncols)

    def header(self, title, subtitle):
        ws = self.ws
        ws["A1"] = title; ws["A1"].font = F_TITLE
        ws["A2"] = subtitle; ws["A2"].font = F_HINT
        ws.cell(3, 1, "Line item").font = F_HDR; ws.cell(3, 1).fill = FILL_HDR
        ws.cell(3, 2, "Where to find it / how to fill").font = F_HDR; ws.cell(3, 2).fill = FILL_HDR
        for i, m in enumerate(self.months):
            c = ws.cell(3, FIRST_COL + i, m); c.font = F_HDR; c.fill = FILL_HDR; c.alignment = Alignment(horizontal="center")
        c = ws.cell(3, FIRST_COL + self.ncols, "Total / Avg"); c.font = F_HDR; c.fill = FILL_HDR; c.alignment = Alignment(horizontal="center")
        ws.column_dimensions["A"].width = 58
        ws.column_dimensions["B"].width = 46
        for i in range(self.ncols + 1):
            ws.column_dimensions[L(FIRST_COL + i)].width = 12
        ws.freeze_panes = ws.cell(4, FIRST_COL)
        ws.sheet_view.zoomScale = 90

    def put(self, line, ct=None, block=""):
        ws, r = self.ws, self.row
        kind = line["kind"]
        key = line["key"]
        ct = ct or {}
        label = resolve(line.get("label", ""), ct)
        if kind == "blank":
            self.row += 1; return
        if kind == "header":
            c = ws.cell(r, 1, label); c.font = F_BOLD; c.fill = FILL_SECTION
            for i in range(1, self.ncols + 1):
                ws.cell(r, FIRST_COL + i - 1).fill = FILL_SECTION
            ws.cell(r, 2).fill = FILL_SECTION; ws.cell(r, FIRST_COL + self.ncols).fill = FILL_SECTION
            self.row += 1; return
        self.rows[key] = r
        lc = ws.cell(r, 1, label); lc.font = F_BOLD if line.get("bold") else F_TEXT
        where = resolve(line.get("where", ""), ct)
        fmt = FMT.get(line.get("fmt"), self.mfmt) if kind != "pct" else FMT["pct"]
        if kind in ("input", "input_first"):
            ws.cell(r, 2, where).font = F_HINT
            ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
            n = self.ncols if kind == "input" else 1
            for i in range(n):
                c = ws.cell(r, FIRST_COL + i); c.font = F_INPUT; c.fill = FILL_INPUT; c.number_format = fmt; c.border = BOX
            if kind == "input_first":
                roll = line.get("roll_from")
                for i in range(1, self.ncols):
                    self.pending.append((ws.cell(r, FIRST_COL + i), ("ROLL", roll, i), F_TEXT, fmt))
            self.inputs.append((r, label))
            self._total(r, kind, line, fmt)
            note = self._note(line, ct)
            if note:
                lc.comment = Comment(note, "d2c-mis"); lc.comment.width = 420; lc.comment.height = 220
        elif kind == "link":
            sheet, k = line["link"]
            for i in range(self.ncols):
                self.pending.append((ws.cell(r, FIRST_COL + i), ("LINK", sheet, k, i, line.get("negate", False)), F_LINK, fmt))
            self._total(r, "formula", line, fmt)
            ws.cell(r, 2, "Linked from another tab; do not type here.").font = F_HINT
        else:  # formula / pct
            ws.cell(r, 2, "Formula; do not type here." if kind == "formula" else "Ratio; do not type here.").font = F_HINT
            for i in range(self.ncols):
                self.pending.append((ws.cell(r, FIRST_COL + i), (line, i), F_BOLD if line.get("bold") else F_TEXT, fmt))
            self._total(r, kind, line, fmt)
            if line.get("bold"):
                for i in range(-2, self.ncols + 1):
                    ws.cell(r, FIRST_COL + i).fill = FILL_SUB
        self.defs.append((block, label, resolve(line.get("what", ""), ct), where, resolve(line.get("mistake", ""), ct)))
        self.row += 1

    def _total(self, r, kind, line, fmt):
        c = self.ws.cell(r, FIRST_COL + self.ncols)
        font = F_BOLD if line.get("bold") else F_TEXT
        c.number_format = fmt; c.font = font
        rng = f"{self.col(0)}{r}:{self.col(self.ncols-1)}{r}"
        if kind == "pct" or line.get("fmt") in ("mult", "days") or line["key"] in ("aov",):
            self.pending.append((c, ("TOTALRATIO", line), font, fmt))
        elif line.get("balance"):
            c.value = f"={self.col(self.ncols-1)}{r}"
        else:
            c.value = f"=SUM({rng})"

    def _note(self, line, ct):
        parts = []
        what = resolve(line.get("what", ""), ct)
        where = resolve(line.get("where", ""), ct)
        mistake = resolve(line.get("mistake", ""), ct)
        if what: parts.append("WHAT: " + what)
        if where: parts.append("WHERE: " + where)
        if mistake: parts.append("COMMON MISTAKE: " + mistake)
        return "\n\n".join(parts)

    def ref(self, key, i):
        """Same-sheet cell reference for key in month i, or '0' if the line is absent."""
        r = self.rows.get(key)
        return f"{self.col(i)}{r}" if r else "0"

    def render(self, extra_refs=None):
        """Resolve pending formulas. extra_refs: dict key -> callable(i) -> reference string."""
        extra_refs = extra_refs or {}
        for cell, spec, font, fmt in self.pending:
            cell.font = font; cell.number_format = fmt
            if isinstance(spec, str):
                cell.value = spec; continue
            if spec[0] == "ROLL":
                _, roll_key, i = spec
                cell.value = f"={self.col(i-1)}{self.rows[roll_key]}"; continue
            if spec[0] == "LINK":
                _, sheet, k, i, negate = spec
                target = extra_refs["__link__"](sheet, k, i)
                cell.value = f"={'-' if negate else ''}{target}"; continue
            if spec[0] == "TOTALRATIO":
                line = spec[1]
                cell.value = self._formula(line, self.ncols, total=True, extra_refs=extra_refs); continue
            line, i = spec
            cell.value = self._formula(line, i, extra_refs=extra_refs)

    def _formula(self, line, i, total=False, extra_refs=None):
        tpl = line["formula"]
        if tpl.startswith("MOM:"):
            k = tpl[4:].strip("{}")
            if total or i == 0: return None
            a, b = self.ref(k, i), self.ref(k, i - 1)
            return f"=IF({b}=0,\"\",{a}/{b}-1)"
        if tpl.startswith("PREV:"):
            k = tpl[5:].strip("{}")
            if total: return f"=SUM({self.col(0)}{self.rows[line['key']]}:{self.col(self.ncols-1)}{self.rows[line['key']]})"
            if i == 0: return f"={self.ref(k, 0)}-{extra_refs['__open_nwc__']()}" if "__open_nwc__" in extra_refs else "=0"
            return f"={self.ref(k, i)}-{self.ref(k, i-1)}"
        if tpl == "RUNWAY":
            if total: return None
            if i < 2: return None
            ob = self.rows["op_burn"]; cc = self.ref("close_cash", i)
            avg = f"AVERAGE({self.col(i-2)}{ob}:{self.col(i)}{ob})"
            return f"=IF({avg}>=0,\"\",-{cc}/{avg})"

        def sub(m):
            k = m.group(1)
            if k in extra_refs: return extra_refs[k](i, total)
            if total:
                r = self.rows.get(k)
                return f"{self.total_col}{r}" if r else "0"
            return self.ref(k, i)
        body = re.sub(r"\{(\w+)\}", sub, tpl)
        if total and line["kind"] == "pct" and "nsr" in tpl:
            pass
        guard = line.get("guard")
        if line["kind"] == "pct":
            denom = tpl.split("/")[-1].strip("()")
            g = re.sub(r"\{(\w+)\}", sub, denom)
            return f"=IF({g}=0,\"\",{body})"
        if guard:
            g = re.sub(r"\{(\w+)\}", sub, guard)
            return f"=IF({g}=0,\"\",{body})"
        return "=" + body


def build(cfg, out_path, preserve_path=None):
    cfg.setdefault("months", 12); cfg.setdefault("unit", "INR lakh"); cfg.setdefault("decisions", {})
    arch = ARCHETYPES[cfg["archetype"]]
    for k, d in DECISIONS.items():
        if d["applies_if"](cfg): cfg["decisions"].setdefault(k, d["default"])
    months = month_labels(cfg["start_month"], cfg["months"])
    wb = Workbook(); wb.remove(wb.active)

    # ---------------- Start Here ----------------
    sh = wb.create_sheet(START)
    sh.column_dimensions["A"].width = 44; sh.column_dimensions["B"].width = 26; sh.column_dimensions["C"].width = 70
    sh["A1"] = f"{cfg['brand']} — Monthly MIS"; sh["A1"].font = F_TITLE
    sh["A2"] = "Investor-grade P&L, channel P&Ls, working capital and cash. Generated by the d2c-mis skill."; sh["A2"].font = F_HINT
    r = 4
    sh.cell(r, 1, "HOW TO USE THIS WORKBOOK").font = F_BOLD; sh.cell(r, 1).fill = FILL_SECTION; r += 1
    for t in [
        "1. Fill the blue cells on each channel tab once a month, from your dashboards. Hover a line label for what goes there, where to find it, and the common mistake.",
        "2. Fill brand marketing and fixed costs on the Consolidated P&L tab. Everything else there is a formula that sums the channel tabs.",
        "3. Fill month-end inventory, receivables, payables and the cash lines on the Cash & Working Capital tab. Closing cash should match your bank.",
        "4. Look at the Checks tab. Anything marked CHECK is worth a second look before you share.",
        "5. Enter amounts in the unit shown below, excluding GST on costs (you claim the input credit) and including GST on gross sales (the sheet backs it out).",
        "6. Blue cells are inputs. Black cells are formulas; do not overwrite them. Green cells are links from other tabs.",
    ]:
        sh.cell(r, 1, t).font = F_TEXT; sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        sh.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top"); sh.row_dimensions[r].height = 30; r += 1
    r += 1
    sh.cell(r, 1, "SETTINGS").font = F_BOLD; sh.cell(r, 1).fill = FILL_SECTION; r += 1
    settings = [("Brand", cfg["brand"]), ("Category archetype", arch["label"]), ("Unit for all amounts", cfg["unit"]),
                ("First month", months[0]), ("Months in this workbook", cfg["months"]), ("Generated on", datetime.date.today().isoformat())]
    for k, v in settings:
        sh.cell(r, 1, k).font = F_TEXT; c = sh.cell(r, 2, v); c.font = F_INPUT; r += 1
    r += 1
    sh.cell(r, 1, "GST RATE BY CHANNEL (edit if your blended rate differs)").font = F_BOLD; sh.cell(r, 1).fill = FILL_SECTION; r += 1
    sh.cell(r, 1, "Channel").font = F_BOLD; sh.cell(r, 2, "Output GST rate").font = F_BOLD; sh.cell(r, 3, "Note").font = F_BOLD; r += 1
    gst_cells = {}
    for ch in cfg["channels"]:
        rate = cfg.get("gst_overrides", {}).get(ch["name"], arch["gst"])
        sh.cell(r, 1, ch["name"]).font = F_TEXT
        c = sh.cell(r, 2, rate); c.font = F_INPUT; c.fill = FILL_KEY; c.number_format = "0.0%"
        sh.cell(r, 3, "Blended rate on your net revenue in this channel. Apparel: 5% under ₹1,000 MRP, 12% above. Food: 5/12/18% by product.").font = F_HINT
        gst_cells[ch["name"]] = f"'{START}'!$B${r}"; r += 1
    r += 1
    sh.cell(r, 1, "PLACEMENT DECISIONS (recorded so anyone reading the MIS knows the convention)").font = F_BOLD; sh.cell(r, 1).fill = FILL_SECTION; r += 1
    sh.cell(r, 1, "Decision").font = F_BOLD; sh.cell(r, 2, "Chosen").font = F_BOLD; sh.cell(r, 3, "Why it matters").font = F_BOLD; r += 1
    for k, d in DECISIONS.items():
        if not d["applies_if"](cfg): continue
        choice = cfg["decisions"][k]
        sh.cell(r, 1, d["question"]).font = F_TEXT; sh.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        sh.cell(r, 2, d["options"][choice]).font = F_TEXT; sh.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        sh.cell(r, 3, d["why"]).font = F_HINT; sh.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
        sh.row_dimensions[r].height = 60; r += 1
    r += 1
    sh.cell(r, 1, "CHANNELS IN THIS WORKBOOK").font = F_BOLD; sh.cell(r, 1).fill = FILL_SECTION; r += 1
    for ch in cfg["channels"]:
        sh.cell(r, 1, ch["name"]).font = F_TEXT; sh.cell(r, 2, CHANNEL_TYPES[ch["type"]]["label"]).font = F_HINT; r += 1
    r += 1
    sh.cell(r, 1, "ASK THE SKILL (things it can help you think through)").font = F_BOLD; sh.cell(r, 1).fill = FILL_SECTION; r += 1
    for q in starter_questions(cfg):
        sh.cell(r, 1, "• " + q).font = F_TEXT; sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); r += 1
    sh.cell(r, 1, "Hover any line label on the other tabs for what goes there. Ask the skill about any line, any ratio, or what an investor would ask. To add a channel later: \"add Amazon\". Your numbers are kept.").font = F_HINT
    sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); sh.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top"); sh.row_dimensions[r].height = 30

    # ---------------- Channel tabs ----------------
    channel_sheets = []
    for ch in cfg["channels"]:
        ct = CHANNEL_TYPES[ch["type"]]
        s = Sheet(wb, ch["name"], cfg, months)
        s.header(f"{ch['name']} — channel P&L", f"{ct['label']}. Blue cells are yours to fill each month, in {cfg['unit']}. Hover a label for guidance.")
        block = ""
        for line in CHANNEL_LINES:
            if line["kind"] == "header": block = line["label"]
            applies = line.get("applies", "all")
            if applies != "all" and ch["type"] not in applies: continue
            cond = line.get("cond")
            if cond and not cond(cfg, ch): continue
            if line["key"] == "cogs":
                line = dict(line, where=arch["cogs_note"])
            s.put(line, ct, block)
        # nsr_base: nsr_post if present else nsr
        base_key = "nsr_post" if "nsr_post" in s.rows else "nsr"
        s.render(extra_refs={
            "gst_rate": lambda i, total, g=gst_cells[ch["name"]]: g,
            "nsr_base": lambda i, total, k=base_key, s=s: (f"{s.total_col}{s.rows[k]}" if total else s.ref(k, i)),
        })
        channel_sheets.append((ch, s))

    # ---------------- Consolidated ----------------
    cs = Sheet(wb, CONS, cfg, months)
    cs.header(f"{cfg['brand']} — Consolidated P&L", f"Channel tabs roll up here automatically. Fill only the blue cells (brand marketing, fixed costs, below-EBITDA). {cfg['unit']}.")
    block = ""
    for key, override in CONSOLIDATED_SUM_KEYS:
        base = next(l for l in CHANNEL_LINES if l["key"] == key)
        if base["kind"] == "header": block = base["label"]
        present = [s for _, s in channel_sheets if key in s.rows]
        if base["kind"] in ("header", "blank"):
            cs.put(base, {}, block); continue
        if not present: continue
        line = dict(base, kind="sum" if base["kind"] == "input" or key in ("gst",) else base["kind"])
        if override:
            line["label"] = override
        elif callable(base.get("label")):
            first_type = next(ch["type"] for ch, s in channel_sheets if s is present[0])
            line["label"] = resolve(base["label"], CHANNEL_TYPES[first_type])
        cs.put_sum(line, present, block) if line["kind"] == "sum" else cs.put(line, {}, block)
    block = "AFTER CM2"
    for line in CONSOLIDATED_LINES:
        if line["kind"] == "header": block = line["label"]
        cs.put(line, {}, block)
    # channel mix rows
    for ch, s in channel_sheets:
        cs.put(dict(key=f"mix_{ch['name']}", kind="pct", label=f"Channel mix: {ch['name']} % of net sales", formula=f"XMIX:{ch['name']}"), {}, "KEY RATIOS & MIX")
    cs_base = "nsr_post" if "nsr_post" in cs.rows else "nsr"
    cs.render(extra_refs={"__xmix__": channel_sheets,
                          "nsr_base": lambda i, total, k=cs_base: (f"{cs.total_col}{cs.rows[k]}" if total else cs.ref(k, i))})

    # ---------------- Cash & WC ----------------
    ca = Sheet(wb, CASH, cfg, months)
    ca.header(f"{cfg['brand']} — Cash & Working Capital", f"Closing balances at month end and the cash roll-forward. {cfg['unit']}. Closing cash should match your bank statement.")
    block = ""
    for line in CASH_LINES:
        if line["kind"] == "header": block = line["label"]
        if line["key"] in ("inventory", "receivables", "payables", "nwc", "open_cash", "close_cash"): line = dict(line, balance=True)
        ca.put(line, {}, block)

    def link(sheet, key, i):
        return f"'{CONS}'!{cs.col(i)}{cs.rows[key]}"
    ca.render(extra_refs={
        "__link__": link,
        "x_cogs": lambda i, total: f"'{CONS}'!{cs.total_col if total else cs.col(i)}{cs.rows['cogs']}",
        "x_nsr": lambda i, total: f"'{CONS}'!{cs.total_col if total else cs.col(i)}{cs.rows['nsr']}",
    })
    # first-month ΔNWC has no prior month: set to 0 with a hint
    ca.ws.cell(ca.rows["d_nwc"], FIRST_COL).value = 0
    ca.ws.cell(ca.rows["d_nwc"], 2).value = "First month is 0 (no prior balance). Later months = this month's NWC less last month's."
    # balance-type totals show last month
    for k in ("inventory", "receivables", "payables", "nwc", "open_cash", "close_cash", "dio", "dso", "dpo", "ccc", "runway"):
        c = ca.ws.cell(ca.rows[k], FIRST_COL + ca.ncols); c.value = f"={ca.col(ca.ncols-1)}{ca.rows[k]}"

    # ---------------- Checks ----------------
    ck = wb.create_sheet(CHECKS)
    ck["A1"] = "Checks — anything marked CHECK deserves a look before you share this MIS"; ck["A1"].font = F_TITLE
    ck["A2"] = f"Bands are for the {arch['label']} archetype. They are guides, not rules; explain an exception rather than hide it."; ck["A2"].font = F_HINT
    ck.cell(3, 1, "Check").font = F_HDR; ck.cell(3, 1).fill = FILL_HDR; ck.cell(3, 2, "Rule").font = F_HDR; ck.cell(3, 2).fill = FILL_HDR
    for i, m in enumerate(months):
        c = ck.cell(3, FIRST_COL + i, m); c.font = F_HDR; c.fill = FILL_HDR; c.alignment = Alignment(horizontal="center")
    ck.column_dimensions["A"].width = 52; ck.column_dimensions["B"].width = 40
    for i in range(len(months)): ck.column_dimensions[L(FIRST_COL + i)].width = 11
    ck.freeze_panes = ck.cell(4, FIRST_COL)
    b = arch["bands"]
    r = 4
    def ck_row(label, rule, fn):
        nonlocal r
        ck.cell(r, 1, label).font = F_TEXT; ck.cell(r, 2, rule).font = F_HINT
        for i in range(len(months)):
            c = ck.cell(r, FIRST_COL + i, fn(i)); c.font = F_TEXT; c.alignment = Alignment(horizontal="center")
        r += 1
    C = lambda k, i: f"'{CONS}'!{cs.col(i)}{cs.rows[k]}"
    K = lambda k, i: f"'{CASH}'!{ca.col(i)}{ca.rows[k]}"
    def band(expr, lo, hi, active):
        return f'=IF({active}=0,"",IF(AND({expr}>={lo},{expr}<={hi}),"OK","CHECK"))'
    ck_row("Gross margin % within archetype band", f"{b['gm'][0]:.0%}–{b['gm'][1]:.0%} of net sales", lambda i: band(f"{C('gm',i)}/{C('nsr',i)}", *b["gm"], C("nsr", i)))
    ret_expr = lambda i: "(" + "+".join(C(k, i) for k in ("returns", "rto", "returns_refund", "exchanges") if k in cs.rows) + ")"
    if any(k in cs.rows for k in ("returns", "rto")):
        ck_row("Returns & RTO % of GMV within band", f"{b['returns_pct_gmv'][0]:.0%}–{b['returns_pct_gmv'][1]:.0%} of GMV", lambda i: band(f"{ret_expr(i)}/{C('gmv',i)}", *b["returns_pct_gmv"], C("gmv", i)))
    ck_row("Logistics % of net sales within band", f"{b['logistics_pct'][0]:.0%}–{b['logistics_pct'][1]:.0%}", lambda i: band(f"({C('logistics',i)}+{C('warehousing',i)})/{C('nsr',i)}", *b["logistics_pct"], C("nsr", i)))
    ck_row("COGS entered when there are sales", "GMV > 0 but COGS blank", lambda i: f'=IF({C("gmv",i)}=0,"",IF({C("cogs",i)}=0,"CHECK","OK"))')
    ck_row("Marketing entered when there are sales", "GMV > 0 but no marketing anywhere", lambda i: f'=IF({C("gmv",i)}=0,"",IF(({C("perf_mkt",i)}+{C("brand_mkt",i)})=0,"CHECK","OK"))')
    ck_row("Blended RoAS above 1.0x", "Net sales ÷ all marketing ≥ 1.0", lambda i: f'=IF(({C("perf_mkt",i)}+{C("brand_mkt",i)})=0,"",IF({C("nsr",i)}/({C("perf_mkt",i)}+{C("brand_mkt",i)})>=1,"OK","CHECK"))')
    ck_row("CM1 not above gross margin", "Fulfilment costs entered", lambda i: f'=IF({C("nsr",i)}=0,"",IF({C("cm1",i)}<={C("gm",i)},"OK","CHECK"))')
    ck_row("Fixed costs entered when there are sales", "People + tech + admin > 0", lambda i: f'=IF({C("gmv",i)}=0,"",IF(({C("people",i)}+{C("tech",i)}+{C("rent_admin",i)})=0,"CHECK","OK"))')
    ck_row("Inventory days within band", f"{b['dio'][0]}–{b['dio'][1]} days", lambda i: f'=IF({K("inventory",i)}=0,"",IF(AND({K("dio",i)}>={b["dio"][0]},{K("dio",i)}<={b["dio"][1]}),"OK","CHECK"))')
    ck_row("Receivable days within band", f"{b['dso'][0]}–{b['dso'][1]} days", lambda i: f'=IF({K("receivables",i)}=0,"",IF(AND({K("dso",i)}>={b["dso"][0]},{K("dso",i)}<={b["dso"][1]}),"OK","CHECK"))')
    ck_row("Working capital entered", "Inventory or receivables filled when there are sales", lambda i: f'=IF({C("gmv",i)}=0,"",IF(({K("inventory",i)}+{K("receivables",i)})=0,"CHECK","OK"))')
    ck_row("Cash roll-forward populated", "Opening cash present", lambda i: f'=IF({C("gmv",i)}=0,"",IF({K("open_cash",i)}=0,"CHECK","OK"))')
    rng = f"{L(FIRST_COL)}4:{L(FIRST_COL+len(months)-1)}{r-1}"
    ck.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"CHECK"'], fill=PatternFill("solid", fgColor="F8CBAD")))
    ck.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", fgColor="C6EFCE")))

    # ---------------- Definitions ----------------
    df = wb.create_sheet(DEFS)
    df["A1"] = "Definitions — every line, what it means, where to find it, and the common mistake"; df["A1"].font = F_TITLE
    hdrs = ["Tab", "Block", "Line", "What it is", "Where to find it", "Common mistake"]
    for j, h in enumerate(hdrs, 1):
        c = df.cell(3, j, h); c.font = F_HDR; c.fill = FILL_HDR
    widths = [22, 24, 46, 60, 60, 60]
    for j, w in enumerate(widths, 1): df.column_dimensions[L(j)].width = w
    r = 4
    seen = set()
    for title, s in [(ch["name"], s) for ch, s in channel_sheets] + [(CONS, cs), (CASH, ca)]:
        for blk, label, what, where, mistake in s.defs:
            if not what and not where: continue
            for j, v in enumerate([title, blk, label, what, where, mistake], 1):
                c = df.cell(r, j, v); c.font = F_TEXT; c.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
    df.freeze_panes = "A4"

    # ---------------- Preserve ----------------
    if preserve_path:
        n = preserve(wb, preserve_path)
        print(f"preserved {n} input values from {preserve_path}")

    order = [START, CONS] + [ch['name'] for ch in cfg['channels']] + [CASH, CHECKS, DEFS]
    wb._sheets = [wb[t] for t in order]
    wb.save(out_path)
    print(f"saved {out_path}: tabs = {wb.sheetnames}")
    return wb


# ---- Consolidated sum helper bound onto Sheet ----
def _put_sum(self, line, present, block):
    ws, r = self.ws, self.row
    self.rows[line["key"]] = r
    label = line["label"] if isinstance(line["label"], str) else line["key"]
    lc = ws.cell(r, 1, label); lc.font = F_TEXT
    ws.cell(r, 2, "Sum of channel tabs; do not type here.").font = F_HINT
    fmt = FMT.get(line.get("fmt"), self.mfmt)
    for i in range(self.ncols):
        parts = [f"'{s.title}'!{s.col(i)}{s.rows[line['key']]}" for s in present]
        c = ws.cell(r, FIRST_COL + i, "=" + "+".join(parts)); c.font = F_LINK; c.number_format = fmt
    self._total(r, "formula", line, fmt)
    self.defs.append((block, label, line.get("what", "") if isinstance(line.get("what"), str) else "", "", ""))
    self.row += 1
Sheet.put_sum = _put_sum

# ---- channel-mix formula hook ----
_orig_formula = Sheet._formula
def _formula_with_mix(self, line, i, total=False, extra_refs=None):
    tpl = line["formula"]
    if tpl.startswith("XMIX:"):
        name = tpl[5:]
        s = next(s for ch, s in extra_refs["__xmix__"] if ch["name"] == name)
        num = f"'{s.title}'!{s.total_col if total else s.col(i)}{s.rows['nsr']}"
        den = f"{self.total_col if total else self.col(i)}{self.rows['nsr']}"
        return f"=IF({den}=0,\"\",{num}/{den})"
    return _orig_formula(self, line, i, total, extra_refs)
Sheet._formula = _formula_with_mix


def preserve(wb, old_path):
    """Copy input values from an older workbook by (sheet, label, month)."""
    old = load_workbook(old_path)
    n = 0
    for ws in wb.worksheets:
        if ws.title not in old.sheetnames: continue
        ows = old[ws.title]
        omonths = {ows.cell(3, c).value: c for c in range(FIRST_COL, ows.max_column + 1) if ows.cell(3, c).value}
        olabels = {ows.cell(r, 1).value: r for r in range(4, ows.max_row + 1) if ows.cell(r, 1).value}
        for r in range(4, ws.max_row + 1):
            label = ws.cell(r, 1).value
            if label not in olabels: continue
            for c in range(FIRST_COL, ws.max_column + 1):
                m = ws.cell(3, c).value
                cell = ws.cell(r, c)
                if cell.font.color is None or cell.font.color.rgb not in ("000000FF", "FF0000FF"): continue  # inputs only
                if m in omonths:
                    v = ows.cell(olabels[label], omonths[m]).value
                    if v is not None and not (isinstance(v, str) and v.startswith("=")):
                        cell.value = v; n += 1
    return n


if __name__ == "__main__":
    args = sys.argv[1:]
    if len(args) < 2:
        print(__doc__); sys.exit(1)
    cfg = json.load(open(args[0]))
    prev = args[args.index("--preserve") + 1] if "--preserve" in args else None
    build(cfg, args[1], prev)
