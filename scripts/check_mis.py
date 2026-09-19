#!/usr/bin/env python3
"""Read a filled MIS workbook and report: which months are filled, the Checks tab results, key ratios,
and anomalies worth asking the founder about. Prints Markdown.

Usage: python check_mis.py filled.xlsx

Values: uses cached values if the file was saved by Excel or Google Sheets. If formulas have no cached
values (file never opened in a spreadsheet app), computes them with the `formulas` library when installed.
"""
import sys, re
from openpyxl import load_workbook

FIRST_COL = 3
CONS, CASH, CHECKS = "Consolidated P&L", "Cash & Working Capital", "Checks"


def load_values(path):
    wb = load_workbook(path, data_only=True)
    ws = wb[CONS]
    probe = [ws.cell(r, FIRST_COL).value for r in range(4, ws.max_row + 1) if (ws.cell(r, 1).value or "").isupper()]
    if any(v is not None for v in probe):
        return wb, lambda sheet, ref: wb[sheet][ref].value
    try:
        import formulas
    except ImportError:
        print("This file has never been opened in Excel or Google Sheets, so formulas have no values yet. "
              "Open it, let it calculate, save, and run again (or `pip install formulas`).")
        sys.exit(1)
    xl = formulas.ExcelModel().loads(path).finish(); sol = xl.calculate()
    fname = path.split("/")[-1]
    def get(sheet, ref):
        v = sol.get(f"'[{fname}]{sheet.upper()}'!{ref}")
        if v is None:
            return wb[sheet][ref].value
        v = v.value[0][0] if hasattr(v, "value") else v
        return None if isinstance(v, str) and v == "" else v
    return wb, get


def rows_by_label(ws):
    return {ws.cell(r, 1).value: r for r in range(4, ws.max_row + 1) if ws.cell(r, 1).value}


def main():
    path = sys.argv[1]
    wb, get = load_values(path)
    cons = wb[CONS]; months = [cons.cell(3, c).value for c in range(FIRST_COL, cons.max_column) if cons.cell(3, c).value and cons.cell(3, c).value != "Total / Avg"]
    cols = [cons.cell(3, c).column_letter for c in range(FIRST_COL, FIRST_COL + len(months))]
    R = rows_by_label(cons)
    unit = wb["Start Here"]["B" + str(next(r for r in range(1, 40) if wb["Start Here"].cell(r, 1).value == "Unit for all amounts"))].value

    def v(label, col, sheet=CONS, rows=None):
        rows = rows or R
        r = rows.get(label)
        if not r: return None
        x = get(sheet, f"{col}{r}")
        try:
            return None if x in (None, "") or isinstance(x, str) else float(x)
        except (TypeError, ValueError):
            return None

    gmv_label = next((l for l in R if l.startswith("Gross sales")), None)
    filled = [i for i, c in enumerate(cols) if (v(gmv_label, c) or 0) > 0]
    print(f"# MIS check — {wb['Start Here']['A1'].value}\n")
    if not filled:
        print("No month has gross sales yet. Fill at least one month on the channel tabs, then run again."); return
    print(f"Months with data: {', '.join(months[i] for i in filled)}. Amounts in {unit}.\n")

    # Key ratios table for filled months
    print("## Key numbers\n")
    keys = [("Net sales", "NET SALES (post GST)", "money"), ("Gross margin %", "Gross margin %", "pct"), ("CM1 %", "CM1 %", "pct"),
            ("CM2 %", "CM2 %", "pct"), ("CM3 %", "CM3 %", "pct"), ("Operating EBITDA", "OPERATING EBITDA", "money"),
            ("Operating EBITDA %", "Operating EBITDA %", "pct"), ("Blended RoAS", "Blended RoAS (net sales ÷ all marketing)", "mult"),
            ("Net sales growth MoM", "Net sales growth month on month", "pct")]
    print("| | " + " | ".join(months[i] for i in filled) + " |"); print("|---|" + "---|" * len(filled))
    def fmt(x, kind):
        if x is None: return "–"
        return {"money": f"{x:,.2f}", "pct": f"{x:.1%}", "mult": f"{x:.1f}x"}[kind]
    for name, label, kind in keys:
        print(f"| {name} | " + " | ".join(fmt(v(label, cols[i]), kind) for i in filled) + " |")
    ca = wb[CASH]; RC = rows_by_label(ca)
    for name, label, kind in [("Closing cash", "CLOSING CASH", "money"), ("Operating burn", "Operating burn (before equity & debt)", "money"),
                              ("Inventory days", "Inventory days (on COGS)", "money"), ("Receivable days", "Receivable days (on net sales)", "money"),
                              ("Cash conversion cycle", "Cash conversion cycle (days)", "money")]:
        print(f"| {name} | " + " | ".join(fmt(v(label, cols[i], CASH, RC), kind) for i in filled) + " |")

    # Checks tab
    ck = wb[CHECKS]; flags = []
    for r in range(4, ck.max_row + 1):
        lab = ck.cell(r, 1).value
        if not lab: continue
        bad = [months[i] for i in filled if get(CHECKS, f"{cols[i]}{r}") == "CHECK"]
        if bad: flags.append(f"- **{lab}** — {', '.join(bad)}. Rule: {ck.cell(r, 2).value}")
    print("\n## Checks\n")
    print("\n".join(flags) if flags else "All checks pass for the filled months.")

    # Anomalies beyond the Checks tab
    print("\n## Worth asking about\n")
    notes = []
    for a, b in zip(filled, filled[1:]):
        g = v("Net sales growth month on month", cols[b])
        if g is not None and abs(g) > 0.5: notes.append(f"- Net sales moved {g:+.0%} from {months[a]} to {months[b]}. Real, or a missing channel / month?")
    for ch in [s for s in wb.sheetnames if s not in ("Start Here", CONS, CASH, CHECKS, "Definitions")]:
        RR = rows_by_label(wb[ch])
        for i in filled:
            cm1 = v("CM1 %", cols[i], ch, RR)
            if cm1 is not None and cm1 < -0.005: notes.append(f"- {ch} CM1 is negative in {months[i]} ({cm1:.0%}). Selling below the cost of delivering.")
            roas = v("RoAS (net sales ÷ marketing)", cols[i], ch, RR)
            if roas is not None and roas < 1.5 and ch.lower().startswith("web"): notes.append(f"- {ch} RoAS {roas:.1f}x in {months[i]}: below 1.5x on the website is hard to sustain.")
    last = filled[-1]
    cc = v("CLOSING CASH", cols[last], CASH, RC); ob = v("Operating burn (before equity & debt)", cols[last], CASH, RC)
    if cc is not None and ob is not None and ob < 0: notes.append(f"- At {months[last]}'s operating burn of {abs(ob):,.2f}, closing cash of {cc:,.2f} is about {cc/abs(ob):.0f} months of runway.")
    print("\n".join(dict.fromkeys(notes)) if notes else "Nothing beyond the Checks tab.")


if __name__ == "__main__":
    main()
