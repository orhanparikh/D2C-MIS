#!/usr/bin/env python3
"""Fill every input cell of a generated MIS with deterministic test values, compute all formulas with the
`formulas` library, and report errors plus a few key rows. Development tool; founders never run this.

Usage: python verify_mis.py workbook.xlsx [--keep filled.xlsx]
"""
import sys, os, random, re, tempfile, shutil
from openpyxl import load_workbook

FIRST_COL = 3
INPUT_RGB = {"000000FF", "FF0000FF"}


def fill_inputs(path, out):
    wb = load_workbook(path)
    rng = random.Random(7)
    n = 0
    for ws in wb.worksheets:
        if ws.title in ("Start Here", "Checks", "Definitions"): continue
        for r in range(4, ws.max_row + 1):
            label = (ws.cell(r, 1).value or "").lower()
            for c in range(FIRST_COL, ws.max_column + 1):
                cell = ws.cell(r, c)
                if cell.font.color is None or cell.font.color.rgb not in INPUT_RGB: continue
                if cell.value not in (None, ""): continue
                # scale by line so ratios look sane: gmv ~ 40, costs fractions of it
                if "gross sales" in label: v = rng.uniform(30, 60)
                elif "orders shipped" in label: v = rng.randint(1500, 3500)
                elif "opening cash" in label: v = 250
                elif "inventory" in label: v = rng.uniform(40, 90)
                elif "receivables" in label: v = rng.uniform(8, 20)
                elif "payables" in label: v = rng.uniform(15, 40)
                elif "equity" in label: v = 0
                elif "debt" in label: v = 0
                elif "other cash" in label: v = 0
                elif "depreciation" in label: v = 0.3
                elif "interest" in label: v = 0.2
                elif "one-time" in label: v = 0
                elif "cost of goods" in label: v = rng.uniform(9, 14)
                elif "gst" in label: v = 0
                else: v = round(rng.uniform(0.3, 6), 2)
                cell.value = round(v, 2); n += 1
    wb.save(out)
    return n


def compute(path):
    import formulas
    xl = formulas.ExcelModel().loads(path).finish()
    sol = xl.calculate()
    out = {}
    for k, v in sol.items():
        m = re.match(r"'\[(.+?)\](.+?)'!(\w+)", k)
        if not m: continue
        sheet, ref = m.group(2), m.group(3)
        val = v.value[0][0] if hasattr(v, "value") else v
        out[(sheet.upper(), ref)] = val
    return out


def main():
    src = sys.argv[1]
    keep = sys.argv[sys.argv.index("--keep") + 1] if "--keep" in sys.argv else None
    tmp = keep or os.path.join(tempfile.mkdtemp(), "filled.xlsx")
    n = fill_inputs(src, tmp)
    print(f"filled {n} input cells -> {tmp}")
    vals = compute(tmp)
    errors = {k: v for k, v in vals.items() if (isinstance(v, str) and v.startswith("#")) or (hasattr(v, "dtype") and str(v).startswith("#"))}
    print(f"computed {len(vals)} cells; errors: {len(errors)}")
    for k, v in list(errors.items())[:30]:
        print("  ERR", k, v)
    # key rows
    wb = load_workbook(tmp)
    for title in wb.sheetnames:
        if title in ("Start Here", "Definitions"): continue
        ws = wb[title]
        print(f"\n== {title}")
        for r in range(4, ws.max_row + 1):
            lab = ws.cell(r, 1).value
            if not lab: continue
            if title == "Checks" or lab.isupper() or lab.endswith("%") or "days" in lab.lower() or "RoAS" in lab or "Runway" in lab or "mix" in lab.lower():
                cells = [vals.get((title.upper(), f"{ws.cell(3, c).column_letter}{r}")) for c in range(FIRST_COL, min(ws.max_column, FIRST_COL + 4))]
                fmt = lambda x: ("" if x in (None, "") else (f"{x:,.2f}" if isinstance(x, (int, float)) else str(x)))
                print(f"  {lab[:55]:55} " + " | ".join(f"{fmt(x):>10}" for x in cells))


if __name__ == "__main__":
    main()
